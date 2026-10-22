import csv
import json

from openpyxl import load_workbook

import sshawk
from sshawk.models import DeviceInfo, Neighbor, StackMember, Status


def sample():
    ok = DeviceInfo(
        host="10.0.0.1",
        status=Status.OK,
        vendor="Cisco",
        os="IOS-XE",
        hostname="SW-CORE-01",
        model="C9300-48P",
        version="17.09.04a",
        serial="FOC1",
        credential="tacacs",
    )
    ok.members = [
        StackMember("1", "active", "C9300-48P", "FOC1", "17.09.04a"),
        StackMember("2", "member", "C9300-48P", "FOC2", "17.09.04a"),
    ]
    ok.neighbors = [Neighbor("cdp", "Te1/1/1", "SW-ACC-01", "Te1/0/1", "10.0.0.2")]
    odd = DeviceInfo(
        host="10.0.0.3",
        status=Status.OK,
        vendor="Cisco",
        os="IOS-XE",
        hostname="SW-ÄUSSERE",
        model="C9300-48P",
        version="17.06.05",
        serial="FOC3",
    )
    failed = DeviceInfo(
        host="10.0.0.9", status=Status.AUTH_FAILED, error="authentication failed (tacacs)"
    )
    return [ok, odd, failed]


def test_json_roundtrip(tmp_path):
    path = tmp_path / "out.json"
    sshawk.export(sample(), path)
    assert (
        json.loads(path.read_text(encoding="utf-8"))["devices"][0]["members"][1]["serial"] == "FOC2"
    )
    loaded = sshawk.load_json(path)
    assert loaded[0].members[1].serial == "FOC2"
    assert loaded[2].status == Status.AUTH_FAILED
    assert loaded[0].neighbors[0].remote_ip == "10.0.0.2"


def test_csv_writes_side_tables(tmp_path):
    written = sshawk.export(sample(), tmp_path / "inv.csv", csv_delimiter=";")
    assert [p.name for p in written] == ["inv.csv", "inv_members.csv", "inv_neighbors.csv"]
    raw = written[0].read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    rows = list(csv.reader(written[0].read_text(encoding="utf-8-sig").splitlines(), delimiter=";"))
    assert rows[0][:3] == ["Host", "Status", "Hostname"]
    assert rows[1][7] == "FOC1, FOC2"
    assert rows[2][2] == "SW-ÄUSSERE"


def test_xlsx_sheets(tmp_path):
    path = tmp_path / "inv.xlsx"
    sshawk.export(sample(), path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Devices", "Stack members", "Neighbors"]
    assert workbook["Devices"]["C2"].value == "SW-CORE-01"
    assert workbook["Stack members"].max_row == 3  # header + 2 members of the stack


def test_html_report(tmp_path):
    path = tmp_path / "inv.html"
    sshawk.export(sample(), path, title="Kunde AG")
    html = path.read_text(encoding="utf-8")
    assert "<title>Kunde AG</title>" in html
    assert "2 releases in use" in html
    assert 'class="port base"' in html and 'class="port s1"' in html
    assert "Login failed (1)" in html
    assert "http" not in html.replace("http://www.w3.org", "")  # no external resources
