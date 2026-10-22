"""Excel export: one sheet per table, with filters, frozen header and sized columns."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from sshawk.export import (
    DEVICE_COLUMNS,
    MEMBER_COLUMNS,
    NEIGHBOR_COLUMNS,
    device_rows,
    member_rows,
    neighbor_rows,
    titles,
)
from sshawk.models import DeviceInfo

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_STATUS_FILLS = {
    "ok": None,
    "partial": PatternFill("solid", fgColor="FFF2CC"),
}
_FAILED_FILL = PatternFill("solid", fgColor="F8D7DA")


def write_xlsx(devices: Sequence[DeviceInfo], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Devices"
    _fill(sheet, titles(DEVICE_COLUMNS), device_rows(devices))
    status_col = titles(DEVICE_COLUMNS).index("Status") + 1
    for row in sheet.iter_rows(min_row=2):
        status = row[status_col - 1].value
        fill = _STATUS_FILLS.get(status, _FAILED_FILL)
        if fill:
            for cell in row:
                cell.fill = fill

    if rows := member_rows(devices):
        _fill(workbook.create_sheet("Stack members"), titles(MEMBER_COLUMNS), rows)
    if rows := neighbor_rows(devices):
        _fill(workbook.create_sheet("Neighbors"), titles(NEIGHBOR_COLUMNS), rows)
    workbook.save(path)


def _fill(sheet: Any, header: list[str], rows: list[list[Any]]) -> None:
    sheet.append(header)
    for cell in sheet[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, title in enumerate(header, start=1):
        values = [title] + [str(r[index - 1]) for r in rows[:500] if r[index - 1] is not None]
        width = min(max(len(v) for v in values) + 2, 60)
        sheet.column_dimensions[get_column_letter(index)].width = width
